from django.core.management.base import BaseCommand
import openpyxl
from epcdata.motorpartsdata.models import Part, PricingData
from oscar.apps.catalogue.models import Product
from oscar.apps.partner.models import StockRecord
from decimal import Decimal
import os

class Command(BaseCommand):
    help = 'Update parts prices from an Excel file (with enhanced debugging).'

    def add_arguments(self, parser):
        parser.add_argument(
            '--verbose',
            action='store_true',
            help='Show detailed output'
        )
        parser.add_argument(
            '--debug',
            action='store_true',
            help='Show extensive debugging information'
        )
        parser.add_argument(
            '--sync-to-oscar',
            action='store_true',
            help='Also update Django Oscar StockRecord prices'
        )

    def handle(self, *args, **options):
        verbose = options.get('verbose', False)
        debug = options.get('debug', False)
        sync_to_oscar = options.get('sync_to_oscar', False)
        
        file_path = os.path.join('epcdata', 'PRCJUL25.xlsx')
        
        self.stdout.write("="*80)
        self.stdout.write("ENHANCED PRICE UPDATE COMMAND (DEBUG VERSION)")
        self.stdout.write("="*80)
        self.stdout.write(f"File path: {file_path}")
        self.stdout.write(f"Verbose mode: {verbose}")
        self.stdout.write(f"Debug mode: {debug}")
        self.stdout.write(f"Sync to Oscar: {sync_to_oscar}")
        self.stdout.write("")

        # Check if file exists
        if not os.path.exists(file_path):
            self.stdout.write(self.style.ERROR(f"❌ File not found: {file_path}"))
            
            # Look for alternative files
            search_dir = 'epcdata'
            if os.path.exists(search_dir):
                excel_files = [f for f in os.listdir(search_dir) if f.endswith(('.xlsx', '.xls'))]
                if excel_files:
                    self.stdout.write(f"📁 Found Excel files in {search_dir}:")
                    for f in excel_files:
                        self.stdout.write(f"  - {f}")
                else:
                    self.stdout.write(f"📁 No Excel files found in {search_dir}")
            return

        self.stdout.write(f"📂 Opening workbook: {file_path}")

        try:
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active
            self.stdout.write(f"✅ Workbook loaded successfully.")
            self.stdout.write(f"📊 Sheet name: {sheet.title}")
            self.stdout.write(f"📊 Sheet dimensions: {sheet.max_row} rows x {sheet.max_column} columns")
            
            if debug:
                # Show first few rows
                self.stdout.write("\n📋 First 5 rows of data:")
                for i, row in enumerate(sheet.iter_rows(min_row=1, max_row=5, values_only=True), 1):
                    self.stdout.write(f"  Row {i}: {row}")
                self.stdout.write("")
                
        except FileNotFoundError:
            self.stdout.write(self.style.ERROR(f"❌ File not found: {file_path}"))
            return
        except Exception as e:
            self.stdout.write(self.style.ERROR(f"❌ Error opening workbook: {e}"))
            return

        # Statistics tracking
        stats = {
            'rows_processed': 0,
            'pricing_data_updated': 0,
            'oscar_updated': 0,
            'parts_not_found': 0,
            'oscar_not_found': 0,
            'errors': 0,
        }

        self.stdout.write("🔄 Processing rows...")
        
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            sku = row[0]
            price = row[1]

            if not sku:
                if debug:
                    self.stdout.write(f"Row {row_num}: ⚠️  Empty SKU, skipping")
                continue

            stats['rows_processed'] += 1
            
            if debug:
                self.stdout.write(f"\n📋 Row {row_num}: Processing SKU '{sku}', Price: {price}")

            try:
                # Update PricingData model
                part = Part.objects.get(part_number=sku)
                pricing_data, created = PricingData.objects.get_or_create(part_number=part)
                
                old_price = pricing_data.list_price
                pricing_data.list_price = price
                pricing_data.price_updated = True
                pricing_data.save()
                
                stats['pricing_data_updated'] += 1
                action = "Created" if created else "Updated"
                
                if verbose or debug:
                    self.stdout.write(f"  ✅ {action} PricingData: {old_price} → {price}")

                # Also update Django Oscar StockRecord if requested
                if sync_to_oscar:
                    try:
                        # Find matching Oscar product
                        oscar_product = Product.objects.filter(upc=sku).first()
                        if oscar_product:
                            stock_record = StockRecord.objects.filter(product=oscar_product).first()
                            if stock_record:
                                old_oscar_price = stock_record.price
                                stock_record.price = Decimal(str(price))
                                stock_record.save()
                                
                                # Verify save
                                stock_record.refresh_from_db()
                                if stock_record.price == Decimal(str(price)):
                                    stats['oscar_updated'] += 1
                                    if verbose or debug:
                                        self.stdout.write(f"  ✅ Updated Oscar StockRecord: £{old_oscar_price} → £{stock_record.price}")
                                else:
                                    self.stdout.write(f"  ❌ Oscar save verification failed for {sku}")
                            else:
                                stats['oscar_not_found'] += 1
                                if debug:
                                    self.stdout.write(f"  ⚠️  No Oscar StockRecord found for {sku}")
                        else:
                            stats['oscar_not_found'] += 1
                            if debug:
                                self.stdout.write(f"  ⚠️  No Oscar Product found for UPC {sku}")
                    except Exception as oscar_error:
                        self.stdout.write(f"  ❌ Oscar update error for {sku}: {oscar_error}")

                if not debug and stats['rows_processed'] % 50 == 0:
                    self.stdout.write(f"  Processed {stats['rows_processed']} rows...")

            except Part.DoesNotExist:
                stats['parts_not_found'] += 1
                if verbose or debug:
                    self.stdout.write(f"  ⚠️  Part not found: {sku}")
            except Exception as e:
                stats['errors'] += 1
                self.stdout.write(f"  ❌ Error processing {sku}: {e}")

        # Final statistics
        self.stdout.write("\n" + "="*60)
        self.stdout.write("ENHANCED SUMMARY")
        self.stdout.write("="*60)
        self.stdout.write(f"File: {file_path}")
        self.stdout.write(f"Rows processed: {stats['rows_processed']}")
        self.stdout.write(f"PricingData updated: {stats['pricing_data_updated']}")
        if sync_to_oscar:
            self.stdout.write(f"Oscar StockRecords updated: {stats['oscar_updated']}")
            self.stdout.write(f"Oscar products not found: {stats['oscar_not_found']}")
        self.stdout.write(f"Parts not found: {stats['parts_not_found']}")
        self.stdout.write(f"Errors: {stats['errors']}")
        
        if not sync_to_oscar:
            self.stdout.write("")
            self.stdout.write("⚠️  NOTE: Only PricingData was updated, not Oscar StockRecords!")
            self.stdout.write("   Use --sync-to-oscar flag to also update prices displayed on website")
            self.stdout.write("")
            self.stdout.write("   To sync prices to Oscar:")
            self.stdout.write("   python manage.py update_prices --sync-to-oscar")

        self.stdout.write(f"\n✅ Price update complete!")
